#!/usr/bin/env python3
"""Read a bounded Excel range or rows matching header filters."""

from __future__ import annotations

import argparse
import sys

from common import add_output_argument, emit, json_value, require_path


def parse_filters(items: list[str]) -> dict[str, str]:
    result = {}
    for item in items:
        if "=" not in item:
            raise ValueError(f"Invalid filter {item!r}; use HEADER=VALUE")
        key, value = item.split("=", 1)
        result[key.strip()] = value.casefold()
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("document")
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--range", dest="cell_range", help="Excel range such as A1:H40")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--filter", action="append", default=[], metavar="HEADER=VALUE")
    parser.add_argument("--max-rows", type=int, default=100)
    parser.add_argument("--max-columns", type=int, default=50)
    parser.add_argument("--data-only", action="store_true")
    add_output_argument(parser)
    args = parser.parse_args()
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries

        path = require_path(args.document, {".xlsx", ".xlsm"})
        if args.max_rows < 1 or args.max_columns < 1 or args.header_row < 1:
            raise ValueError("row and column limits must be positive")
        workbook = load_workbook(path, read_only=True, data_only=args.data_only, keep_links=False)
        if args.sheet not in workbook.sheetnames:
            raise ValueError(f"Unknown sheet {args.sheet!r}; choose from {workbook.sheetnames}")
        sheet = workbook[args.sheet]
        if args.cell_range:
            min_col, min_row, max_col, max_row = range_boundaries(args.cell_range)
        else:
            min_col, min_row = 1, args.header_row
            max_col, max_row = sheet.max_column or 1, sheet.max_row or args.header_row
        max_col = min(max_col, min_col + args.max_columns - 1)
        filters = parse_filters(args.filter)
        header_values = [cell.value for cell in next(sheet.iter_rows(
            min_row=args.header_row, max_row=args.header_row, min_col=min_col, max_col=max_col
        ))]
        header_index = {str(value).strip(): index for index, value in enumerate(header_values) if value is not None}
        missing = sorted(set(filters) - set(header_index))
        if missing:
            raise ValueError(f"Filter headers not found: {missing}")
        rows = []
        scanned = 0
        truncated = False
        for cells in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            scanned += 1
            values = [json_value(cell.value) for cell in cells]
            if filters and cells[0].row != args.header_row:
                if not all(filters[key] in str(values[header_index[key]]).casefold() for key in filters):
                    continue
            rows.append({"row": cells[0].row, "cells": values})
            if len(rows) >= args.max_rows:
                truncated = cells[0].row < max_row
                break
        end_row = rows[-1]["row"] if rows else None
        emit({
            "document": str(path), "type": "excel", "sheet": sheet.title,
            "source": {"range": args.cell_range, "start_row": min_row, "end_row": end_row},
            "headers": [json_value(value) for value in header_values], "rows": rows,
            "scanned_rows": scanned, "truncated": truncated,
        }, args.output)
        workbook.close()
        return 0
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

