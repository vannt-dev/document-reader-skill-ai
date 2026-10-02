#!/usr/bin/env python3
"""Create a bounded, deterministic index for a local requirements folder."""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from common import DOCUMENT_TYPES, add_output_argument, emit, json_value

SUPPORTED = set(DOCUMENT_TYPES)
HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")


def digest(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def markdown_metadata(path: Path, max_headings: int) -> dict:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    headings = []
    fenced = False
    for number, line in enumerate(lines, 1):
        stripped = line.lstrip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            fenced = not fenced
        if not fenced and (match := HEADING.match(line)):
            headings.append({
                "level": len(match.group(1)),
                "title": match.group(2).rstrip("# "),
                "line": number,
            })
            if len(headings) > max_headings:
                break
    return {
        "line_count": len(lines),
        "headings": headings[:max_headings],
        "headings_truncated": len(headings) > max_headings,
    }


def excel_metadata(path: Path, preview_rows: int, max_columns: int) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel indexing requires openpyxl: python -m pip install openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            preview = []
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row or 1, preview_rows),
                max_col=min(sheet.max_column or 1, max_columns),
                values_only=True,
            ):
                preview.append([json_value(value) for value in row])
            sheets.append({
                "name": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "preview": preview,
                "preview_truncated_columns": (sheet.max_column or 0) > max_columns,
            })
    finally:
        workbook.close()
    return {"sheets": sheets}


def pdf_metadata(path: Path, max_headings: int) -> dict:
    from doc_text import open_pdf, pdf_bookmarks

    reader = open_pdf(path)
    bookmarks, truncated = pdf_bookmarks(reader, max_headings)
    return {
        "page_count": len(reader.pages),
        "bookmarks": bookmarks,
        "bookmarks_truncated": truncated,
    }


def docx_metadata(path: Path, max_headings: int) -> dict:
    from doc_text import docx_items, docx_outline

    items = docx_items(path)
    headings = docx_outline(items)
    return {
        "paragraph_count": len(items),
        "table_rows": sum(1 for item in items if item["kind"] == "table_row"),
        "headings": headings[:max_headings],
        "headings_truncated": len(headings) > max_headings,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", help="Authorized requirements folder")
    parser.add_argument("--recursive", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--max-files", type=int, default=500)
    parser.add_argument("--preview-rows", type=int, default=3)
    parser.add_argument("--max-columns", type=int, default=20)
    parser.add_argument("--max-headings", type=int, default=200)
    add_output_argument(parser)
    args = parser.parse_args()
    try:
        root = Path(args.folder).expanduser().resolve()
        if not root.is_dir():
            raise FileNotFoundError(f"Requirements folder not found: {root}")
        if min(args.max_files, args.preview_rows, args.max_columns, args.max_headings) < 1:
            raise ValueError("All limits must be positive")
        iterator = root.rglob("*") if args.recursive else root.glob("*")
        candidates = sorted(
            (path for path in iterator if path.is_file() and path.suffix.lower() in SUPPORTED),
            key=lambda item: item.relative_to(root).as_posix().casefold(),
        )
        truncated = len(candidates) > args.max_files
        documents = []
        for path in candidates[:args.max_files]:
            stat = path.stat()
            item = {
                "path": path.relative_to(root).as_posix(),
                "type": DOCUMENT_TYPES[path.suffix.lower()],
                "size_bytes": stat.st_size,
                "modified_utc": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                "sha256": digest(path),
            }
            if item["type"] == "markdown":
                item["structure"] = markdown_metadata(path, args.max_headings)
            elif item["type"] in {"pdf", "docx"}:
                # A folder that only held Excel and Markdown before may now pick up a PDF
                # nobody can open (no pypdf, a password, a damaged file). Record that on
                # the document instead of failing the whole index.
                try:
                    metadata = pdf_metadata if item["type"] == "pdf" else docx_metadata
                    item["structure"] = metadata(path, args.max_headings)
                except Exception as exc:  # noqa: BLE001
                    item["structure"] = {"error": str(exc)}
            else:
                item["structure"] = excel_metadata(path, args.preview_rows, args.max_columns)
            documents.append(item)
        emit({
            "schema_version": "1.0",
            "requirements_root": str(root),
            "generated_utc": datetime.now(timezone.utc).isoformat(),
            "documents": documents,
            "truncated": truncated,
            "supported_extensions": sorted(SUPPORTED),
        }, args.output)
        return 0
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
