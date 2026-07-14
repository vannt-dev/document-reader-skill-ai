#!/usr/bin/env python3
"""Inspect workbook structure while returning only small previews."""

from __future__ import annotations

import argparse
import sys

from common import add_output_argument, emit, json_value, require_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("document")
    parser.add_argument("--preview-rows", type=int, default=5)
    parser.add_argument("--max-columns", type=int, default=30)
    parser.add_argument("--data-only", action="store_true")
    add_output_argument(parser)
    args = parser.parse_args()
    try:
        from openpyxl import load_workbook

        path = require_path(args.document, {".xlsx", ".xlsm"})
        if args.preview_rows < 0 or args.max_columns < 1:
            raise ValueError("preview-rows must be >= 0 and max-columns must be >= 1")
        workbook = load_workbook(path, read_only=True, data_only=args.data_only, keep_links=False)
        sheets = []
        for sheet in workbook.worksheets:
            preview = []
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row or 1, args.preview_rows),
                max_col=min(sheet.max_column or 1, args.max_columns),
                values_only=True,
            ):
                preview.append([json_value(value) for value in row])
            sheets.append({
                "name": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "preview": preview,
                "preview_truncated_columns": (sheet.max_column or 0) > args.max_columns,
            })
        emit({"document": str(path), "type": "excel", "data_only": args.data_only, "sheets": sheets}, args.output)
        workbook.close()
        return 0
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

