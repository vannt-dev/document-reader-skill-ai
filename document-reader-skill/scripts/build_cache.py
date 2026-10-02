#!/usr/bin/env python3
"""Extract full document content into a durable per-document JSON cache."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from cache_common import column_letter, doc_id_for, excel_block_ranges, markdown_sections, slugify
from common import add_output_argument, emit, json_value
import cache_config


def extract_markdown(path: Path) -> list[dict]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    blocks = []
    for section in markdown_sections(lines):
        title = section["title"] or "(preamble)"
        text = "\n".join(lines[section["start"] - 1:section["end"]])
        blocks.append({
            "block_id": f"{slugify(title)}@L{section['start']}-{section['end']}",
            "ref": {
                "heading": section["title"],
                "level": section["level"],
                "lines": [section["start"], section["end"]],
            },
            "label": title,
            "content": {"text": text},
        })
    return blocks


def extract_pdf(path: Path) -> list[dict]:
    from doc_text import open_pdf, pdf_page_text

    reader = open_pdf(path)
    blocks = []
    for number in range(1, len(reader.pages) + 1):
        text = pdf_page_text(reader, number)
        if not text:
            continue
        blocks.append({
            "block_id": f"page-{number}",
            "ref": {"page": number},
            "label": f"Page {number}",
            "content": {"text": text},
        })
    if not blocks:
        raise ValueError("No extractable text; a scanned PDF needs OCR before it can be read")
    return blocks + pdf_table_blocks(path)


def pdf_table_blocks(path: Path) -> list[dict]:
    """Tables as their own blocks, shaped like an Excel sheet, so a query can
    return a table with its header instead of a page of loose text.

    Table extraction is optional: without pdfplumber, or when it cannot read
    the file, the text blocks stand on their own.
    """
    import pdf_tables

    if not pdf_tables.available():
        return []
    try:
        tables = pdf_tables.extract_tables(path)
    except Exception:  # noqa: BLE001 - the page text is already cached
        return []
    blocks = []
    for table in tables:
        page, number, rows = table["page"], table["table"], table["rows"]
        blocks.append({
            "block_id": f"page-{page}-table-{number}",
            "ref": {"page": page, "table": number},
            "label": f"Page {page} table {number}",
            "content": {
                "headers": rows[0],
                "rows": [{"row": index, "cells": cells} for index, cells in enumerate(rows[1:], 2)],
            },
        })
    return blocks


def extract_docx(path: Path) -> list[dict]:
    from doc_text import docx_items, docx_render, docx_sections

    items = docx_items(path)
    blocks = []
    for section in docx_sections(items):
        title = section["title"] or "(preamble)"
        blocks.append({
            "block_id": f"{slugify(title)}@P{section['start']}-{section['end']}",
            "ref": {
                "heading": section["title"],
                "level": section["level"],
                "paragraphs": [section["start"], section["end"]],
            },
            "label": title,
            "content": {"text": docx_render(items[section["start"] - 1:section["end"]])},
        })
    return blocks


def extract_excel(path: Path, window: int) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    blocks = []
    try:
        for sheet in workbook.worksheets:
            max_row = sheet.max_row or 0
            max_col = sheet.max_column or 0
            if max_row < 1 or max_col < 1:
                continue
            header_cells = next(
                sheet.iter_rows(min_row=1, max_row=1, max_col=max_col, values_only=True), ()
            )
            headers = [json_value(value) for value in header_cells]
            last_col = column_letter(max_col)
            ranges = excel_block_ranges(max_row, window)
            for start, end in ranges:
                rows = []
                # Row 1 is the header; it lives only in `headers`, never in `rows`,
                # so every block (including the first) carries data rows only.
                data_start = max(start, 2)
                for cells in sheet.iter_rows(min_row=data_start, max_row=end, max_col=max_col):
                    rows.append({
                        "row": cells[0].row,
                        "cells": [json_value(cell.value) for cell in cells],
                    })
                cell_range = f"A{start}:{last_col}{end}"
                label = sheet.title if len(ranges) == 1 else f"{sheet.title} rows {start}-{end}"
                blocks.append({
                    "block_id": f"{sheet.title}!{cell_range}",
                    "ref": {"sheet": sheet.title, "range": cell_range},
                    "label": label,
                    "content": {"headers": headers, "rows": rows},
                })
    finally:
        workbook.close()
    return blocks


def build(manifest_path: Path, cache_dir: Path, window: int, force: bool = False) -> dict:
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    root = Path(manifest["requirements_root"])
    cache_subdir = cache_dir / "cache"
    cache_subdir.mkdir(parents=True, exist_ok=True)
    documents = []
    used_doc_ids: set[str] = set()
    for entry in manifest.get("documents", []):
        rel = entry["path"]
        doc_id = doc_id_for(rel)
        # Different paths can slugify to the same doc_id (e.g. "a-b.md" and
        # "a_b.md"); disambiguate so one cache file never overwrites another.
        if doc_id in used_doc_ids:
            suffix = 2
            while f"{doc_id}-{suffix}" in used_doc_ids:
                suffix += 1
            doc_id = f"{doc_id}-{suffix}"
        used_doc_ids.add(doc_id)
        sha256 = entry.get("sha256", "")
        doc_type = entry.get("type", "excel")
        cache_file = cache_subdir / f"{doc_id}.json"

        if not force and cache_file.is_file():
            existing = json.loads(cache_file.read_text(encoding="utf-8"))
            if existing.get("sha256") == sha256:
                documents.append(_index_entry(existing))
                continue

        cache_payload = {"doc_id": doc_id, "path": rel, "sha256": sha256, "type": doc_type}
        try:
            source = root / rel
            if doc_type == "markdown":
                cache_payload["blocks"] = extract_markdown(source)
            elif doc_type == "pdf":
                cache_payload["blocks"] = extract_pdf(source)
            elif doc_type == "docx":
                cache_payload["blocks"] = extract_docx(source)
            else:
                cache_payload["blocks"] = extract_excel(source, window)
        except Exception as exc:  # noqa: BLE001 - record, do not abort the whole build
            cache_payload["blocks"] = []
            cache_payload["error"] = str(exc)

        cache_file.write_text(
            json.dumps(cache_payload, ensure_ascii=False, indent=2, default=str) + "\n",
            encoding="utf-8",
        )
        documents.append(_index_entry(cache_payload))

    index = {"schema_version": "1.0", "documents": documents}
    (cache_dir / "index.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    return index


def _index_entry(cache_payload: dict) -> dict:
    entry = {
        "doc_id": cache_payload["doc_id"],
        "path": cache_payload["path"],
        "sha256": cache_payload.get("sha256", ""),
        "type": cache_payload.get("type", "unknown"),
        "blocks": [
            {"block_id": block["block_id"], "label": block["label"]}
            for block in cache_payload.get("blocks", [])
        ],
    }
    if "error" in cache_payload:
        entry["error"] = cache_payload["error"]
    return entry


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", default=".document-reader/manifest.json")
    parser.add_argument("--cache-dir", default=".document-reader")
    parser.add_argument("--max-block-rows", type=int, default=500)
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--share", choices=["local", "commit"], help="Set cache share scope and adjust .gitignore")
    add_output_argument(parser)
    args = parser.parse_args()
    try:
        if args.max_block_rows < 1:
            raise ValueError("max-block-rows must be positive")
        manifest_path = Path(args.manifest)
        if not manifest_path.is_file():
            raise FileNotFoundError(f"Manifest not found: {manifest_path}; run index_documents.py first")
        index = build(manifest_path, Path(args.cache_dir), args.max_block_rows, args.force)
        if args.share:
            cache_dir = Path(args.cache_dir)
            cache_config.write_config(cache_dir, args.share)
            cache_config.apply_gitignore(cache_dir.parent, args.share)
            if args.share == "commit":
                print("warning: share=commit will version cache content; verify documents contain no secrets.", file=sys.stderr)
        emit({
            "cache_dir": args.cache_dir,
            "documents": len(index["documents"]),
            "errors": [d["doc_id"] for d in index["documents"] if "error" in d],
        }, args.output)
        return 0
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
