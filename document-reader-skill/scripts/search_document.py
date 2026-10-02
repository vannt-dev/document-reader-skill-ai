#!/usr/bin/env python3
"""Search a document and return compact cited matches.

Excel and CSV cite cells, Markdown cites lines, PDF pages, Word paragraphs and
PowerPoint slides.
"""

from __future__ import annotations

import argparse
import re
import sys

from common import DOCUMENT_TYPES, add_output_argument, emit, json_value, require_path


def matcher(query: str, regex: bool, case_sensitive: bool):
    flags = 0 if case_sensitive else re.IGNORECASE
    pattern = re.compile(query if regex else re.escape(query), flags)
    return pattern.search


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("document")
    parser.add_argument("query")
    parser.add_argument("--sheet", action="append", default=[])
    parser.add_argument("--regex", action="store_true")
    parser.add_argument("--case-sensitive", action="store_true")
    parser.add_argument("--max-results", type=int, default=50)
    parser.add_argument("--context-lines", type=int, default=1)
    add_output_argument(parser)
    args = parser.parse_args()
    try:
        path = require_path(args.document, set(DOCUMENT_TYPES))
        if args.max_results < 1 or args.context_lines < 0:
            raise ValueError("max-results must be positive and context-lines non-negative")
        matches, truncated = [], False
        find = matcher(args.query, args.regex, args.case_sensitive)
        if path.suffix.lower() in {".md", ".markdown"}:
            lines = path.read_text(encoding="utf-8-sig").splitlines()
            hit_lines = [number for number, line in enumerate(lines, 1) if find(line)]
            truncated = len(hit_lines) > args.max_results
            for number in hit_lines[:args.max_results]:
                start, end = max(1, number - args.context_lines), min(len(lines), number + args.context_lines)
                matches.append({"line": number, "context_start": start, "context_end": end, "text": "\n".join(lines[start - 1:end])})
            kind = "markdown"
        elif path.suffix.lower() == ".pdf":
            from doc_text import open_pdf, pdf_page_text

            reader = open_pdf(path)
            for page in range(1, len(reader.pages) + 1):
                lines = pdf_page_text(reader, page).splitlines()
                for number, line in enumerate(lines, 1):
                    if not find(line):
                        continue
                    if len(matches) >= args.max_results:
                        truncated = True
                        break
                    start, end = max(1, number - args.context_lines), min(len(lines), number + args.context_lines)
                    matches.append({"page": page, "line": number, "text": "\n".join(lines[start - 1:end])})
                if truncated:
                    break
            kind = "pdf"
        elif path.suffix.lower() == ".docx":
            from doc_text import docx_items, docx_render

            items = docx_items(path)
            hits = [item["index"] for item in items if find(item["text"])]
            truncated = len(hits) > args.max_results
            for index in hits[:args.max_results]:
                start, end = max(1, index - args.context_lines), min(len(items), index + args.context_lines)
                matches.append({
                    "paragraph": index, "context_start": start, "context_end": end,
                    "text": docx_render(items[start - 1:end]),
                })
            kind = "docx"
        elif path.suffix.lower() == ".pptx":
            from doc_text import pptx_render, pptx_slides

            for slide in pptx_slides(path):
                lines = pptx_render(slide).splitlines()
                for number, line in enumerate(lines, 1):
                    if not find(line):
                        continue
                    if len(matches) >= args.max_results:
                        truncated = True
                        break
                    start, end = max(1, number - args.context_lines), min(len(lines), number + args.context_lines)
                    matches.append({
                        "slide": slide["slide"], "title": slide["title"],
                        "line": number, "text": "\n".join(lines[start - 1:end]),
                    })
                if truncated:
                    break
            kind = "pptx"
        elif path.suffix.lower() in {".csv", ".tsv"}:
            from cache_common import column_letter
            from csv_text import csv_rows

            _, rows = csv_rows(path)
            for number, cells in rows:
                for index, value in enumerate(cells, 1):
                    if not find(value):
                        continue
                    if len(matches) >= args.max_results:
                        truncated = True
                        break
                    matches.append({"row": number, "cell": f"{column_letter(index)}{number}", "value": value})
                if truncated:
                    break
            kind = "csv"
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            names = args.sheet or workbook.sheetnames
            unknown = sorted(set(names) - set(workbook.sheetnames))
            if unknown:
                raise ValueError(f"Unknown sheets: {unknown}")
            for name in names:
                for row in workbook[name].iter_rows():
                    for cell in row:
                        value = json_value(cell.value)
                        if value is not None and find(str(value)):
                            matches.append({"sheet": name, "cell": cell.coordinate, "value": value})
                            if len(matches) >= args.max_results:
                                truncated = True
                                break
                    if truncated:
                        break
                if truncated:
                    break
            workbook.close()
            kind = "excel"
        emit({"document": str(path), "type": kind, "query": args.query, "matches": matches, "truncated": truncated}, args.output)
        return 0
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

